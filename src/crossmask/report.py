from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import RunResult


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _format_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for column_cells in sheet.columns:
        length = max((len(str(cell.value or "")) for cell in column_cells), default=0)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(length + 2, 12), 42)


def write_report(result: RunResult) -> Path:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "Value"])
    summary.append(["Generated at (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    summary.append(["Files processed", len(result.processed_files)])
    summary.append(["Cells transformed", sum(item.cells_changed for item in result.transformations)])
    summary.append(["Residual findings", len(result.verification_findings)])
    summary.append(["Raw values stored in report", "No"])
    summary.append(["Processing mode", "Local only"])
    _format_sheet(summary)

    transformations = workbook.create_sheet("Transformations")
    transformations.append(["File", "Sheet", "Column", "Entity", "Action", "Cells changed"])
    for item in result.transformations:
        transformations.append([item.file, item.sheet, item.column, item.entity, item.action, item.cells_changed])
    _format_sheet(transformations)

    verification = workbook.create_sheet("Verification")
    verification.append(["File", "Sheet", "Row", "Column", "Entity", "Value fingerprint"])
    for item in result.verification_findings:
        verification.append([item.file, item.sheet, item.row, item.column, item.entity, item.fingerprint])
    _format_sheet(verification)

    result.report_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(result.report_path)
    return result.report_path
