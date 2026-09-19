from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .models import Table


SUPPORTED_SUFFIXES = {".csv", ".xlsx"}


def iter_supported_files(source: str | Path) -> list[Path]:
    path = Path(source)
    if path.is_file():
        if path.suffix.casefold() not in SUPPORTED_SUFFIXES:
            raise ValueError(f"Unsupported file type: {path.suffix}")
        return [path]
    if not path.is_dir():
        raise FileNotFoundError(f"Input does not exist: {path}")
    return sorted(
        item
        for item in path.rglob("*")
        if item.is_file()
        and item.suffix.casefold() in SUPPORTED_SUFFIXES
        and not item.stem.endswith("_sanitized")
        and item.name != "crossmask_privacy_report.xlsx"
    )


def _csv_encoding(path: Path) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            path.read_text(encoding=encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("crossmask", b"", 0, 1, f"Unable to decode {path}")


def load_tables(path: Path) -> list[Table]:
    if path.suffix.casefold() == ".csv":
        with path.open("r", encoding=_csv_encoding(path), newline="") as handle:
            rows = list(csv.reader(handle))
        if not rows:
            return [Table(name="CSV", headers=[], rows=[])]
        headers = [str(value or f"column_{index + 1}") for index, value in enumerate(rows[0])]
        width = len(headers)
        normalized_rows = [(row + [""] * width)[:width] for row in rows[1:]]
        return [Table(name="CSV", headers=headers, rows=normalized_rows)]

    workbook = load_workbook(path, data_only=False, read_only=False)
    tables: list[Table] = []
    for sheet in workbook.worksheets:
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            tables.append(Table(name=sheet.title, headers=[], rows=[]))
            continue
        headers = [str(value or f"column_{index + 1}") for index, value in enumerate(values[0])]
        width = len(headers)
        rows = [list(row[:width]) + [None] * max(0, width - len(row)) for row in values[1:]]
        tables.append(Table(name=sheet.title, headers=headers, rows=rows))
    return tables


def write_tables(source: Path, tables: list[Table], destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if source.suffix.casefold() == ".csv":
        table = tables[0]
        with destination.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(table.headers)
            writer.writerows(table.rows)
        return

    workbook = Workbook()
    workbook.remove(workbook.active)
    for table in tables:
        sheet = workbook.create_sheet(title=table.name[:31] or "Sheet")
        if table.headers:
            sheet.append(table.headers)
        for row in table.rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        if table.headers:
            sheet.auto_filter.ref = sheet.dimensions
    if not workbook.worksheets:
        workbook.create_sheet("Sheet1")
    workbook.save(destination)
