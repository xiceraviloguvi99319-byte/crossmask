from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from crossmask.engine import run_path, scan_path, verify_path


SECRET = "unit-test-secret-at-least-12"


class CrossMaskTests(unittest.TestCase):
    def test_consistent_aliases_across_csv_files(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            (input_dir / "people.csv").write_text(
                "full_name,phone\nLin Tao,13800138000\nMaya Chen,13900139000\n",
                encoding="utf-8",
            )
            (input_dir / "orders.csv").write_text(
                "order_id,full_name,phone\nO-1,Lin Tao,13800138000\n",
                encoding="utf-8",
            )

            result = run_path(input_dir, output_dir, SECRET)

            with (output_dir / "people_sanitized.csv").open(encoding="utf-8-sig", newline="") as handle:
                people = list(csv.DictReader(handle))
            with (output_dir / "orders_sanitized.csv").open(encoding="utf-8-sig", newline="") as handle:
                orders = list(csv.DictReader(handle))

            self.assertEqual(people[0]["full_name"], orders[0]["full_name"])
            self.assertEqual(people[0]["phone"], orders[0]["phone"])
            self.assertTrue(people[0]["full_name"].startswith("Person-"))
            self.assertEqual([], result.verification_findings)
            self.assertTrue(result.report_path.exists())

    def test_xlsx_pipeline_and_report(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "linked.xlsx"
            output_dir = root / "output"
            workbook = Workbook()
            people = workbook.active
            people.title = "People"
            people.append(["姓名", "身份证号", "邮箱"])
            people.append(["Zhang Wei", "110101199001011234", "zhang.wei@example.test"])
            orders = workbook.create_sheet("Orders")
            orders.append(["order_id", "姓名"])
            orders.append(["O-9", "Zhang Wei"])
            workbook.save(source)

            result = run_path(source, output_dir, SECRET)
            output_book = load_workbook(output_dir / "linked_sanitized.xlsx", data_only=False)

            self.assertEqual(output_book["People"]["A2"].value, output_book["Orders"]["B2"].value)
            self.assertTrue(str(output_book["People"]["B2"].value).startswith("ID-"))
            self.assertTrue(str(output_book["People"]["C2"].value).endswith("@example.invalid"))
            self.assertEqual([], result.verification_findings)

    def test_scan_and_verify(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "raw.csv"
            source.write_text("notes\nCall 13800138000\n", encoding="utf-8")

            scan_items = scan_path(source)
            verify_items = verify_path(source)

            self.assertEqual("phone", scan_items[0].entity)
            self.assertEqual("value pattern", scan_items[0].detection_source)
            self.assertEqual("phone", verify_items[0].entity)
            self.assertNotIn("13800138000", verify_items[0].fingerprint)


if __name__ == "__main__":
    unittest.main()
